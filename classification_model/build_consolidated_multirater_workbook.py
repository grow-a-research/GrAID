"""
build_consolidated_multirater_workbook.py — one workbook, one sheet per
essay (200 sheets: 185 new + the 15 reference), each sheet showing the
essay's ID/Band/Text plus a full 5-professor scoring grid (Content/
Organization/Language, 1-5 each), matching the layout of the original
15-essay multi-rater workbook — this is the "all 5 professors grade all
200 essays" design, not the single-rater 37-per-professor split.

Column layout per sheet (matches the provided example):
  A Essay ID | B Band | C Essay Text
  D-H Content (Prof 1-5) | I Content Avg
  J-N Organization (Prof 1-5) | O Organization Avg
  P-T Language (Prof 1-5) | U Language Avg
  V Total (/15) | W Overall %

Output: classification_model/dress_data/consolidated_multirater.xlsx (gitignored)
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .build_professor_batch2_workbooks import SHARED_15, select_new_185
from .dress_pipeline import load_dress_bulk

OUT_PATH = "classification_model/dress_data/consolidated_multirater.xlsx"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
GROUP_FILL = PatternFill("solid", fgColor="2E75B6")
AVG_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
WHITE_BOLD = Font(name=FONT_NAME, bold=True, color="FFFFFF")
BOLD = Font(name=FONT_NAME, bold=True)
REGULAR = Font(name=FONT_NAME)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)

PROFESSORS = ["Professor 1", "Professor 2", "Professor 3", "Professor 4", "Professor 5"]
CRITERIA = ["CONTENT", "ORGANIZATION", "LANGUAGE"]
GROUP_START_COL = {"CONTENT": "D", "ORGANIZATION": "J", "LANGUAGE": "P"}
AVG_COL = {"CONTENT": "I", "ORGANIZATION": "O", "LANGUAGE": "U"}


def build_essay_sheet(wb, essay_id: str, band: str, essay_text: str):
    ws = wb.create_sheet(essay_id)

    for col, label in [("A", "Essay ID"), ("B", "Band"), ("C", "Essay Text")]:
        ws.merge_cells(f"{col}1:{col}2")
        c = ws[f"{col}1"]
        c.value = label
        c.font = WHITE_BOLD
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws[f"{col}2"].border = BORDER
        ws[f"{col}2"].fill = HEADER_FILL

    for crit in CRITERIA:
        start = GROUP_START_COL[crit]
        start_idx = ws[f"{start}1"].column
        end_letter = get_column_letter(start_idx + 4)
        ws.merge_cells(f"{start}1:{end_letter}1")
        c = ws[f"{start}1"]
        c.value = crit
        c.font = WHITE_BOLD
        c.fill = GROUP_FILL
        c.alignment = CENTER
        c.border = BORDER
        for i, prof in enumerate(PROFESSORS):
            col_letter = get_column_letter(start_idx + i)
            cell = ws[f"{col_letter}2"]
            cell.value = prof
            cell.font = WHITE_BOLD
            cell.fill = GROUP_FILL
            cell.alignment = CENTER
            cell.border = BORDER

        avg_letter = AVG_COL[crit]
        ws.merge_cells(f"{avg_letter}1:{avg_letter}2")
        c = ws[f"{avg_letter}1"]
        c.value = f"{crit.title()}\nAvg"
        c.font = BOLD
        c.fill = AVG_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws[f"{avg_letter}2"].border = BORDER
        ws[f"{avg_letter}2"].fill = AVG_FILL

    for col, label in [("V", "Total\n(/15)"), ("W", "Overall\n%")]:
        ws.merge_cells(f"{col}1:{col}2")
        c = ws[f"{col}1"]
        c.value = label
        c.font = BOLD
        c.fill = AVG_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws[f"{col}2"].border = BORDER
        ws[f"{col}2"].fill = AVG_FILL

    dv = DataValidation(type="whole", operator="between", formula1=1, formula2=5,
                         allow_blank=True, showErrorMessage=True,
                         errorTitle="Invalid score", error="Score must be a whole number from 1 to 5.")
    ws.add_data_validation(dv)

    row = 3
    ws[f"A{row}"] = essay_id
    ws[f"A{row}"].font = BOLD
    ws[f"A{row}"].alignment = CENTER
    ws[f"A{row}"].border = BORDER
    ws[f"B{row}"] = band
    ws[f"B{row}"].font = BOLD
    ws[f"B{row}"].alignment = CENTER
    ws[f"B{row}"].border = BORDER
    ws[f"C{row}"] = essay_text
    ws[f"C{row}"].font = REGULAR
    ws[f"C{row}"].alignment = LEFT
    ws[f"C{row}"].border = BORDER

    for crit in CRITERIA:
        start = GROUP_START_COL[crit]
        start_idx = ws[f"{start}1"].column
        for i in range(5):
            col_letter = get_column_letter(start_idx + i)
            cell = ws[f"{col_letter}{row}"]
            cell.fill = INPUT_FILL
            cell.font = REGULAR
            cell.alignment = CENTER
            cell.border = BORDER
            dv.add(cell)
        avg_letter = AVG_COL[crit]
        end_letter = get_column_letter(start_idx + 4)
        avg_cell = ws[f"{avg_letter}{row}"]
        avg_cell.value = f'=IFERROR(AVERAGE({start}{row}:{end_letter}{row}),"")'
        avg_cell.font = BOLD
        avg_cell.alignment = CENTER
        avg_cell.border = BORDER

    total_cell = ws[f"V{row}"]
    total_cell.value = f'=IFERROR(I{row}+O{row}+U{row},"")'
    total_cell.font = BOLD
    total_cell.alignment = CENTER
    total_cell.border = BORDER

    pct_cell = ws[f"W{row}"]
    pct_cell.value = f'=IFERROR(V{row}/15*100,"")'
    pct_cell.number_format = "0.0"
    pct_cell.font = BOLD
    pct_cell.alignment = CENTER
    pct_cell.border = BORDER

    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 11
    ws.column_dimensions["C"].width = 70
    for col in "DEFGHJKLMNPQRST":
        ws.column_dimensions[col].width = 9
    for col in ["I", "O", "U"]:
        ws.column_dimensions[col].width = 10
    ws.column_dimensions["V"].width = 9
    ws.column_dimensions["W"].width = 9
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 220
    ws.freeze_panes = "D3"


def main():
    bulk = load_dress_bulk()
    new_185 = select_new_185()
    all_ids = SHARED_15 + new_185["essay_id"].tolist()
    essays = bulk[bulk["essay_id"].isin(all_ids)][["essay_id", "essay_text", "gt_band"]]
    print(f"Building {len(essays)} essay sheets (expected 200)...")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    index_ws = wb.create_sheet("Index", 0)
    index_ws.append(["Essay ID", "Band"])
    for c in ("A1", "B1"):
        index_ws[c].font = WHITE_BOLD
        index_ws[c].fill = HEADER_FILL
    index_ws.column_dimensions["A"].width = 14
    index_ws.column_dimensions["B"].width = 12

    for _, row in essays.sort_values("essay_id").iterrows():
        build_essay_sheet(wb, row["essay_id"], row["gt_band"], row["essay_text"])
        index_ws.append([row["essay_id"], row["gt_band"]])

    wb.save(OUT_PATH)
    print(f"Saved {OUT_PATH} ({len(essays)} essay sheets + 1 index sheet)")


if __name__ == "__main__":
    main()
