"""
build_score_entry_workbook.py — builds dress_shared_scores_entry.xlsx, a
human-friendly wide-format spreadsheet for transcribing the 5 professors'
paper scores on the 15-essay DREsS_New shared subset.

This is NOT the format classification_model.load_professor_grades() reads
directly (that needs the long, one-row-per-criterion-per-professor CSV).
This workbook is for easy manual entry; a separate conversion step will
reshape it into that long format once it's filled in.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_PATH = "classification_model/dress_shared_scores_entry.xlsx"

FONT_NAME = "Arial"

ESSAYS = [
    ("DRESS-1790", 168),
    ("DRESS-1781", 360),
    ("DRESS-1779", 276),
    ("DRESS-1777", 361),
    ("DRESS-1432", 226),
    ("DRESS-2151", 411),
    ("DRESS-1450", 385),
    ("DRESS-1685", 403),
    ("DRESS-1778", 347),
    ("DRESS-1449", 335),
    ("DRESS-1412", 444),
    ("DRESS-1664", 328),
    ("DRESS-1652", 391),
    ("DRESS-1433", 408),
    ("DRESS-2193", 236),
]

PROFESSORS = ["Professor 1", "Professor 2", "Professor 3", "Professor 4", "Professor 5"]
CRITERIA = ["CONTENT", "ORGANIZATION", "LANGUAGE"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
GROUP_FILL = PatternFill("solid", fgColor="2E75B6")
AVG_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
WHITE_BOLD = Font(name=FONT_NAME, bold=True, color="FFFFFF")
BOLD = Font(name=FONT_NAME, bold=True)
REGULAR = Font(name=FONT_NAME)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def build_scores_sheet(wb):
    ws = wb.active
    ws.title = "Scores - Fill In"

    # Column layout:
    # A Essay ID | B Word Count
    # C-G Content P1-P5 | H Content Avg
    # I-M Organization P1-P5 | N Organization Avg
    # O-S Language P1-P5 | T Language Avg
    # U Total (/15) | V Overall %
    single_cols = {
        "A": "Essay ID",
        "B": "Word\nCount",
    }
    group_start_col = {"CONTENT": "C", "ORGANIZATION": "I", "LANGUAGE": "O"}
    avg_col = {"CONTENT": "H", "ORGANIZATION": "N", "LANGUAGE": "T"}

    # Row 1+2 headers for single (non-grouped) columns, merged vertically
    for col, label in single_cols.items():
        ws.merge_cells(f"{col}1:{col}2")
        c = ws[f"{col}1"]
        c.value = label
        c.font = WHITE_BOLD
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws[f"{col}2"].border = BORDER
        ws[f"{col}2"].fill = HEADER_FILL

    # Grouped criterion headers (row 1 merged across 5 prof columns, row 2 = prof names)
    for crit in CRITERIA:
        start = group_start_col[crit]
        start_idx = ws[f"{start}1"].column
        end_letter = get_column_letter(start_idx + 4)
        ws.merge_cells(f"{start}1:{end_letter}1")
        c = ws[f"{start}1"]
        c.value = crit
        c.font = WHITE_BOLD
        c.fill = GROUP_FILL
        c.alignment = CENTER
        c.border = BORDER
        for i, prof in enumerate(PROFESSORS):
            col_letter = get_column_letter(start_idx + i)
            cell = ws[f"{col_letter}2"]
            cell.value = prof
            cell.font = WHITE_BOLD
            cell.fill = GROUP_FILL
            cell.alignment = CENTER
            cell.border = BORDER

        avg_letter = avg_col[crit]
        ws.merge_cells(f"{avg_letter}1:{avg_letter}2")
        c = ws[f"{avg_letter}1"]
        c.value = f"{crit.title()}\nAvg"
        c.font = BOLD
        c.fill = AVG_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws[f"{avg_letter}2"].border = BORDER
        ws[f"{avg_letter}2"].fill = AVG_FILL

    for col, label in [("U", "Total\n(/15)"), ("V", "Overall\n%")]:
        ws.merge_cells(f"{col}1:{col}2")
        c = ws[f"{col}1"]
        c.value = label
        c.font = BOLD
        c.fill = AVG_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws[f"{col}2"].border = BORDER
        ws[f"{col}2"].fill = AVG_FILL

    # Data validation: whole number 1-5 for all professor-score input cells
    dv = DataValidation(
        type="whole", operator="between", formula1=1, formula2=5,
        allow_blank=True, showErrorMessage=True,
        errorTitle="Invalid score", error="Score must be a whole number from 1 to 5.",
    )
    ws.add_data_validation(dv)

    first_data_row = 3
    for row_offset, (essay_id, word_count) in enumerate(ESSAYS):
        row = first_data_row + row_offset
        ws[f"A{row}"] = essay_id
        ws[f"A{row}"].font = BOLD
        ws[f"A{row}"].alignment = LEFT
        ws[f"A{row}"].border = BORDER
        ws[f"B{row}"] = word_count
        ws[f"B{row}"].font = REGULAR
        ws[f"B{row}"].alignment = CENTER
        ws[f"B{row}"].border = BORDER

        for crit in CRITERIA:
            start = group_start_col[crit]
            start_idx = ws[f"{start}1"].column
            for i in range(5):
                col_letter = get_column_letter(start_idx + i)
                cell = ws[f"{col_letter}{row}"]
                cell.fill = INPUT_FILL
                cell.font = REGULAR
                cell.alignment = CENTER
                cell.border = BORDER
                dv.add(cell)
            avg_letter = avg_col[crit]
            end_letter = get_column_letter(start_idx + 4)
            avg_cell = ws[f"{avg_letter}{row}"]
            avg_cell.value = f'=IFERROR(AVERAGE({start}{row}:{end_letter}{row}),"")'
            avg_cell.font = BOLD
            avg_cell.alignment = CENTER
            avg_cell.border = BORDER

        total_cell = ws[f"U{row}"]
        total_cell.value = f'=IFERROR(H{row}+N{row}+T{row},"")'
        total_cell.font = BOLD
        total_cell.alignment = CENTER
        total_cell.border = BORDER

        pct_cell = ws[f"V{row}"]
        pct_cell.value = f'=IFERROR(U{row}/15*100,"")'
        pct_cell.number_format = "0.0"
        pct_cell.font = BOLD
        pct_cell.alignment = CENTER
        pct_cell.border = BORDER

    # Column widths
    widths = {"A": 13, "B": 9}
    for col in "CDEFGIJKLMOPQRS":
        widths[col] = 9
    for col in ["H", "N", "T"]:
        widths[col] = 10
    widths["U"] = 9
    widths["V"] = 9
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 22

    ws.freeze_panes = "C3"
    return ws


def build_reference_sheet(wb):
    ws = wb.create_sheet("Essay Reference")
    headers = ["Essay ID", "Prompt", "Word Count"]
    prompt = (
        "If you could change one important thing about your university, what "
        "would you change? Give specific reasons and details to support your opinion."
    )
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = WHITE_BOLD
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
    for row_offset, (essay_id, word_count) in enumerate(ESSAYS):
        row = row_offset + 2
        ws.cell(row=row, column=1, value=essay_id).font = REGULAR
        ws.cell(row=row, column=1).alignment = CENTER
        ws.cell(row=row, column=2, value=prompt).font = REGULAR
        ws.cell(row=row, column=2).alignment = LEFT
        ws.cell(row=row, column=3, value=word_count).font = REGULAR
        ws.cell(row=row, column=3).alignment = CENTER
        for col in range(1, 4):
            ws.cell(row=row, column=col).border = BORDER
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 11
    ws.freeze_panes = "A2"
    note = ws.cell(row=len(ESSAYS) + 3, column=1,
                    value="Same prompt for all 15 essays (this is the shared DREsS_New batch). "
                          "Full essay text is not repeated here to keep this sheet short — see "
                          "dress_shared_subset_candidates-1.pdf for the full text of each essay.")
    note.font = Font(name=FONT_NAME, italic=True, size=9)
    ws.merge_cells(start_row=note.row, start_column=1, end_row=note.row, end_column=3)


def build_instructions_sheet(wb):
    ws = wb.create_sheet("Instructions", 0)
    ws.sheet_view.showGridLines = False
    title = ws.cell(row=1, column=1, value="How to fill in this workbook")
    title.font = Font(name=FONT_NAME, bold=True, size=14)

    lines = [
        "",
        "1. Go to the \"Scores - Fill In\" tab.",
        "2. Only edit the YELLOW cells — one score per essay, per criterion, per professor.",
        "3. Each score is 1-5 (whole numbers), matching GRAID-RUBRICS.pdf. A dropdown/validation "
        "will reject anything outside 1-5.",
        "4. The \"Avg\", \"Total (/15)\", and \"Overall %\" columns are formulas — they fill in "
        "automatically as you enter scores. Do not type into those.",
        "5. \"Professor 1\"-\"Professor 5\" are placeholders. If you want, rename the header cells "
        "(row 2 under each criterion) to the actual professors' initials/codes — just keep the "
        "SAME professor in the SAME column position across all three criteria (Content, "
        "Organization, Language).",
        "6. The \"Essay Reference\" tab lists the shared prompt and word count per essay, matching "
        "the order in \"Scores - Fill In\", in case you need to double check which essay you're "
        "scoring against the paper sheets.",
        "",
        "Before finalizing: confirm with the professors whether any of them used half-point "
        "scores (e.g. 3.5). The rubric distributed to them (GRAID-RUBRICS.pdf) only shows whole "
        "levels (5/4/3/2/1) — if any half-points were actually used on paper, let us know before "
        "this data goes into training, since the current validation only accepts whole numbers.",
        "",
        "This is a manual data-entry workbook only — it is a separate step from "
        "classification_model.py's training format. Once this is filled in, it still needs to be "
        "reshaped into the long-format professor_grades.csv the training script reads "
        "(essay_id, part, professor_id, question_prompt, essay_text, criterion_name, "
        "criterion_max_points, criterion_score) — that conversion is a next step, not done yet.",
    ]
    for i, line in enumerate(lines, start=2):
        cell = ws.cell(row=i, column=1, value=line)
        cell.font = REGULAR
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        ws.row_dimensions[i].height = 30 if len(line) > 90 else (16 if line else 8)
    ws.column_dimensions["A"].width = 14
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 14


def main():
    wb = Workbook()
    build_scores_sheet(wb)
    build_reference_sheet(wb)
    build_instructions_sheet(wb)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
