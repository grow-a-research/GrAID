"""
professor_ground_truth.py — loads real professor-consensus scores from the
filled-in consolidated single-sheet workbook (200 essays: 15 reference +
185 new, each graded by all 5 professors on Content/Organization/Language).

Consensus = average across the 5 professors, per criterion, summed to a
0-15 total then converted to a percentage — same convention used
throughout this project for DREsS's own scores, so the two are directly
comparable.
"""

from __future__ import annotations

import openpyxl
import pandas as pd

from .dress_pipeline import MAX_TOTAL_POINTS, score_to_band

CONTENT_COLS = range(4, 9)       # D-H
ORG_COLS = range(10, 15)         # J-N
LANG_COLS = range(16, 21)        # P-T


def load_professor_consensus(xlsx_path: str) -> pd.DataFrame:
    """Returns essay_id, prof_score_pct, prof_band for all graded rows in the sheet."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Scores - Fill In"]

    rows = []
    for r in range(3, ws.max_row + 1):
        essay_id = ws.cell(row=r, column=1).value
        if not essay_id:
            continue
        content_avg = sum(ws.cell(row=r, column=c).value for c in CONTENT_COLS) / 5
        org_avg = sum(ws.cell(row=r, column=c).value for c in ORG_COLS) / 5
        lang_avg = sum(ws.cell(row=r, column=c).value for c in LANG_COLS) / 5
        total = content_avg + org_avg + lang_avg
        rows.append({"essay_id": essay_id, "prof_total": total})

    df = pd.DataFrame(rows)
    df["prof_score_pct"] = df["prof_total"] / MAX_TOTAL_POINTS * 100
    df["prof_band"] = df["prof_score_pct"].apply(score_to_band)
    return df[["essay_id", "prof_score_pct", "prof_band"]]
