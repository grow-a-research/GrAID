"""
build_consolidated_essays_workbook.py — consolidates the 200-essay training
set (185 new + the 15 already-graded reference essays) into ONE workbook,
one sheet per essay, each showing the DREsS ID, essay text, and band.

"Band" here is the DREsS-derived band (from DREsS's own official score),
the only band classification available for the 185 new essays before
professor grading happens — shown for reference/review, not as a
grading input.

Output: classification_model/dress_data/consolidated_essays.xlsx (gitignored)
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .build_professor_batch2_workbooks import SHARED_15, select_new_185
from .dress_pipeline import load_dress_bulk

OUT_PATH = "classification_model/dress_data/consolidated_essays.xlsx"

FONT_NAME = "Arial"
LABEL_FILL = PatternFill("solid", fgColor="1F4E78")
WHITE_BOLD = Font(name=FONT_NAME, bold=True, color="FFFFFF")
REGULAR = Font(name=FONT_NAME, size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def build_sheet(wb, essay_id: str, band: str, essay_text: str):
    ws = wb.create_sheet(essay_id)
    ws.sheet_view.showGridLines = False

    ws["A1"] = "DREsS ID"
    ws["B1"] = essay_id
    ws["A2"] = "Band"
    ws["B2"] = band
    ws["A3"] = "Essay Text"

    for cell_ref in ("A1", "A2", "A3"):
        c = ws[cell_ref]
        c.font = WHITE_BOLD
        c.fill = LABEL_FILL
        c.border = BORDER
    for cell_ref in ("B1", "B2"):
        c = ws[cell_ref]
        c.font = Font(name=FONT_NAME, bold=True)
        c.border = BORDER

    ws.merge_cells("B3:F40")
    text_cell = ws["B3"]
    text_cell.value = essay_text
    text_cell.font = REGULAR
    text_cell.alignment = LEFT_WRAP
    text_cell.border = BORDER

    ws.column_dimensions["A"].width = 14
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16


def main():
    bulk = load_dress_bulk()
    new_185 = select_new_185()
    all_ids = SHARED_15 + new_185["essay_id"].tolist()

    essays = bulk[bulk["essay_id"].isin(all_ids)][["essay_id", "essay_text", "gt_band"]]
    print(f"Consolidating {len(essays)} essays (expected 200)...")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    index_ws = wb.create_sheet("Index", 0)
    index_ws.append(["Essay ID", "Band"])
    for c in ("A1", "B1"):
        index_ws[c].font = WHITE_BOLD
        index_ws[c].fill = LABEL_FILL

    for _, row in essays.sort_values("essay_id").iterrows():
        build_sheet(wb, row["essay_id"], row["gt_band"], row["essay_text"])
        index_ws.append([row["essay_id"], row["gt_band"]])

    index_ws.column_dimensions["A"].width = 14
    index_ws.column_dimensions["B"].width = 12

    wb.save(OUT_PATH)
    print(f"Saved {OUT_PATH} ({len(essays)} essay sheets + 1 index sheet)")


if __name__ == "__main__":
    main()
