"""
build_final_200_dataset.py — builds the true final 200-essay dataset:
the 172 essays that survived cleaning (the original 200 minus the 28
confirmed-problematic ones: 16 off-topic + 12 unreliable-re-grade), plus
the 28 freshly-graded replacement essays (new prompt-essay pairing, new
independent 5-professor grading, verified clean on data-integrity checks).

Quartile-recalibrated bands are recomputed fresh on this final 200-essay
population, then the classifier is retrained and evaluated exactly as
before, for a clean, final, defensible comparison.

Run from repo root:

    python -m classification_model.build_final_200_dataset
"""

from pathlib import Path

import openpyxl
import pandas as pd

from .classification_model import (
    BAND_ORDER,
    cross_validate,
    evaluate_predictions,
    summarize_formula,
    train_ordinal_classifier,
)

DATA_DIR = Path(__file__).resolve().parent / "dress_data"
REPLACED_28 = [
    "DRESS-1016", "DRESS-1050", "DRESS-1056", "DRESS-1061", "DRESS-1062", "DRESS-1066", "DRESS-1109",
    "DRESS-1126", "DRESS-1163", "DRESS-1231", "DRESS-1235", "DRESS-1401", "DRESS-1731", "DRESS-1987",
    "DRESS-2005", "DRESS-2172", "DRESS-34", "DRESS-353", "DRESS-587", "DRESS-593", "DRESS-631",
    "DRESS-633", "DRESS-634", "DRESS-636", "DRESS-644", "DRESS-766", "DRESS-906", "DRESS-969",
]
OUT_PATH = DATA_DIR / "final_200_combined.csv"


def load_new_28_professor_scores() -> pd.DataFrame:
    wb = openpyxl.load_workbook(DATA_DIR / "replacement_28_workbook.xlsx", data_only=True)
    ws = wb["Scores - Fill In"]
    rows = []
    for r in range(3, ws.max_row + 1):
        eid = ws.cell(row=r, column=1).value
        if not eid:
            continue
        content = [ws.cell(row=r, column=c).value for c in range(4, 9)]
        org = [ws.cell(row=r, column=c).value for c in range(10, 15)]
        lang = [ws.cell(row=r, column=c).value for c in range(16, 21)]
        if any(v is None for v in content + org + lang):
            raise ValueError(f"Missing scores for {eid} in replacement_28_workbook.xlsx")
        total_pct = (sum(content) / 5 + sum(org) / 5 + sum(lang) / 5) / 15 * 100
        rows.append({"essay_id": eid, "prof_score_pct": total_pct})
    return pd.DataFrame(rows)


def main():
    print("Loading the original 200-essay dataset (few-shot, current final)...")
    original = pd.read_csv(DATA_DIR / "fewshot_full_combined.csv").rename(
        columns={"ai_score_pct": "score_pct", "ai_confidence": "confidence", "ai_spread": "spread"}
    )
    surviving_172 = original[~original["essay_id"].isin(REPLACED_28)][
        ["essay_id", "score_pct", "confidence", "spread", "prof_score_pct"]
    ]
    print(f"  Surviving (unreplaced) essays: {len(surviving_172)}")

    print("Loading the 28 replacement essays' AI + professor scores...")
    ai28 = pd.read_csv(DATA_DIR / "replacement_28_ai_progress.csv").rename(
        columns={"ai_score_pct": "score_pct", "ai_confidence": "confidence", "ai_spread": "spread"}
    )
    prof28 = load_new_28_professor_scores()
    new_28 = ai28.merge(prof28, on="essay_id", how="inner")
    if len(new_28) != 28:
        raise ValueError(f"Expected 28 replacement essays, got {len(new_28)} after merge")
    print(f"  Replacement essays merged: {len(new_28)}")

    combined = pd.concat([surviving_172, new_28], ignore_index=True)
    if len(combined) != 200:
        raise ValueError(f"Final dataset should be 200 essays, got {len(combined)}")
    if combined["essay_id"].duplicated().any():
        raise ValueError(f"Duplicate essay_ids: {combined[combined['essay_id'].duplicated()]['essay_id'].tolist()}")

    print("\nRecomputing quartile-based bands on the final 200-essay professor-score distribution...")
    # Round before banding: the two score sources compute the same percentage via
    # different float arithmetic (e.g. 63.99999999999999 vs 64.0), which silently
    # dropped essays sitting exactly on a cutoff into the lower band. A score equal
    # to a cutoff belongs to the higher band.
    rounded_pct = combined["prof_score_pct"].round(6)
    q25, q50, q75 = rounded_pct.quantile([0.25, 0.5, 0.75]).round(6)
    print(f"  Cutoffs: Poor/Fair={q25:.2f}  Fair/Good={q50:.2f}  Good/Excellent={q75:.2f}")

    def band(pct):
        if pct < q25:
            return "Poor"
        if pct < q50:
            return "Fair"
        if pct < q75:
            return "Good"
        return "Excellent"

    combined["gt_band"] = rounded_pct.apply(band)
    combined.to_csv(OUT_PATH, index=False)
    print(f"  Saved -> {OUT_PATH}")

    print(f"\nBand distribution (n={len(combined)}):")
    print(combined["gt_band"].value_counts().reindex(BAND_ORDER))

    corr = combined["score_pct"].corr(combined["prof_score_pct"])
    bias = (combined["score_pct"] - combined["prof_score_pct"]).mean()
    mae = (combined["score_pct"] - combined["prof_score_pct"]).abs().mean()
    print(f"\nCorrelation: {corr:.3f}   Bias: {bias:+.2f}   MAE: {mae:.2f}")

    print("\nBias by band:")
    for b in BAND_ORDER:
        sub = combined[combined["gt_band"] == b]
        print(f"  {b}: n={len(sub)}  bias={(sub['score_pct']-sub['prof_score_pct']).mean():+.2f}")

    print("\n=== Training final ordinal classifier ===")
    _, result = train_ordinal_classifier(combined)
    print(summarize_formula(result))

    print("\n=== 5-fold cross-validation ===")
    cv = cross_validate(combined)
    ev = evaluate_predictions(cv)
    print(f"  Exact-match accuracy: {ev['exact_accuracy']:.2%}")
    print(f"  Adjacent-band agreement: {ev['adjacent_agreement']:.2%}")
    print(ev["confusion_matrix"])

    majority_band = combined["gt_band"].value_counts().idxmax()
    majority_acc = (combined["gt_band"] == majority_band).mean()
    print(f"\n  Majority baseline ('{majority_band}'): {majority_acc:.2%}")

    print("\nPer-band recall:")
    confusion = ev["confusion_matrix"]
    for b in BAND_ORDER:
        total = confusion.loc[b].sum()
        correct = confusion.loc[b, b]
        print(f"  {b}: {correct}/{total} = {correct/total:.1%}" if total else f"  {b}: n/a")


if __name__ == "__main__":
    main()
