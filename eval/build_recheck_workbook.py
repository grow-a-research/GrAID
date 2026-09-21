"""
build_recheck_workbook.py — the blind second pass.

Why this exists
---------------
On the first pass, 40 of 250 judgments disagreed with GrAid. Auditing them
showed the disagreements are not all GrAid's fault: 15 were answers whose
transcription is character-identical to the answer key (the student wrote it
right, the professor marked it wrong), and 3 were answers belonging to a
different question that were marked right. Those look like marking slips, not
system errors.

Correcting them by hand would be fatal to the study — it would mean deciding
which of the professor's judgments to keep based on whether they agreed with
the system. So instead the professor judges a subset again, blind.

How it stays blind
------------------
The file mixes every disagreement together with a matched sample of items where
the professor and GrAid already agreed, shuffled with a fresh seed. Nothing in
it marks which is which, and as on the first pass nothing shows GrAid's
transcription or score. Including the agreements is what protects the result:
re-checking only the items GrAid got wrong would hand the system a free
correction in one direction.

The agreement sample is drawn to match the disagreements question-for-question,
so the mix of questions cannot hint at which group a row belongs to.

What comes out of it
--------------------
  - corrected judgments to use for the final Analysis C
  - intra-rater reliability: how often the same person judges the same answer
    the same way. That is a Chapter 4 result in its own right, and it puts a
    number on the "professor inconsistency" category in Analysis D.

The answer key mapping each item back to its first-pass judgment is written to
a separate file that must NOT be sent to the professor.

Usage
-----
    ./venv/Scripts/python.exe eval/build_recheck_workbook.py
"""
from __future__ import annotations

import argparse
import csv
import json
import random
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from export_identification import (  # noqa: E402
    DEFAULT_DB, DEFAULT_EXAM, OUT_DIR, _answers, _db, item_id,
)
from judgment_sheet import build_workbook, guard_overwrite  # noqa: E402

RECHECK_SEED = 20260922      # deliberately different from the first pass

INTRO = [
    ("Thank you for a second short round. This one is about 80 answers, not 250.", False),
    ("", False),
    ("What you are doing", True),
    ("These are answers from the first round, shown again. Please judge each one fresh, without "
     "trying to remember what you decided last time. Judging a subset twice is how we measure "
     "how consistently answers are marked, which is a normal and expected part of a study like "
     "this — there is no right or wrong in having changed your mind.", False),
    ("", False),
    ("The instructions below are slightly clearer than last time", True),
    ("Two things were easy to read the wrong way on the first round, so they are spelled out "
     "explicitly now. Please read them even though the task is familiar.", False),
    ("", False),
    ("Why you still cannot see the system's answers", True),
    ("Nothing here shows what GrAid read or what score it gave. Your judgment is the standard we "
     "measure the system against, so it has to stay independent. The rows are shuffled again, in "
     "a different order from last time.", False),
]


def build(args) -> None:
    guard_overwrite(OUT_DIR / args.out, args.force)
    # ── first-pass judgments ────────────────────────────────────────────
    wb1 = openpyxl.load_workbook(args.first_pass, data_only=True)
    ws1 = wb1["Judgments"]
    pass1 = {}
    for r in range(5, ws1.max_row + 1):
        iid, verdict = ws1.cell(r, 2).value, ws1.cell(r, 6).value
        if iid and verdict:
            pass1[iid] = verdict
    if not pass1:
        sys.exit(f"no judgments found in {args.first_pass}")

    # ── what GrAid actually decided ─────────────────────────────────────
    con = _db(args.db)
    rows = {item_id(r["submission_id"], r["order_index"]): r for r in _answers(con, args.exam)}

    agree: dict[int, list[str]] = defaultdict(list)
    disagree: dict[int, list[str]] = defaultdict(list)
    for iid, verdict in pass1.items():
        r = rows.get(iid)
        if r is None:
            continue
        credited = (r["ai_score"] or 0) > 0
        bucket = agree if (verdict == "Correct") == credited else disagree
        bucket[r["order_index"]].append(iid)

    n_dis = sum(len(v) for v in disagree.values())

    # ── matched sample of agreements ────────────────────────────────────
    rng = random.Random(RECHECK_SEED)
    picked: list[str] = []
    for oi, dis in disagree.items():
        pool = list(agree.get(oi, []))
        rng.shuffle(pool)
        picked += pool[:len(dis)]

    # top up to the requested size from whatever agreements are left
    remaining = [i for oi, v in agree.items() for i in v if i not in set(picked)]
    rng.shuffle(remaining)
    picked += remaining[:max(0, args.agreements - len(picked))]
    picked = picked[:args.agreements]

    chosen = [i for v in disagree.values() for i in v] + picked
    rng.shuffle(chosen)

    items = [{"iid": i,
              "prompt": rows[i]["prompt"] or "",
              "accepted": rows[i]["correct_answer"] or ""} for i in chosen]

    dest = OUT_DIR / args.out
    missing = build_workbook(items, OUT_DIR / "crops", dest,
                             "Second round — judging GrAid's identification answers", INTRO)

    # ── the private key ─────────────────────────────────────────────────
    dis_set = {i for v in disagree.values() for i in v}
    key_path = OUT_DIR / "recheck_key_DO_NOT_SEND.csv"
    with key_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["item_id", "group", "pass1_judgment", "graid_credited", "graid_ocr",
                    "answer_key"])
        for i in chosen:
            r = rows[i]
            w.writerow([i, "disagreement" if i in dis_set else "agreement", pass1[i],
                        int((r["ai_score"] or 0) > 0), (r["ocr_text"] or "").strip(),
                        r["correct_answer"] or ""])

    print(f"  wrote {dest.relative_to(REPO)}  "
          f"({len(items)} answers, {dest.stat().st_size / 1024 / 1024:.1f} MB)")
    print(f"  wrote {key_path.relative_to(REPO)}")
    if missing:
        print(f"  {len(missing)} crops were missing: {', '.join(missing)}")

    print(f"\n  Composition: {n_dis} disagreements + {len(picked)} agreements = {len(items)}")
    print("  Per question (disagreements / agreements drawn):")
    picked_by_q: dict[int, int] = defaultdict(int)
    for i in picked:
        picked_by_q[rows[i]["order_index"]] += 1
    for oi in sorted(set(list(disagree) + list(picked_by_q))):
        print(f"    Q{oi:02d}  {len(disagree.get(oi, [])):2d} / {picked_by_q.get(oi, 0):2d}")

    manifest = OUT_DIR / "manifest.json"
    if manifest.exists():
        data = json.loads(manifest.read_text())
        data["recheck_workbook"] = {
            "file": args.out, "items": len(items), "disagreements": n_dis,
            "agreements": len(picked), "seed": RECHECK_SEED,
            "first_pass": str(args.first_pass), "key_file": key_path.name,
        }
        manifest.write_text(json.dumps(data, indent=2))

    print(f"\n  SEND: {dest.name}")
    print(f"  DO NOT SEND: {key_path.name} — it says which rows were disputed.")
    print("\n  Note for Chapter 3: the second pass used clarified instructions, so the")
    print("  two passes are not identical conditions. Say so when reporting the")
    print("  intra-rater figure — the clarification was applied to every row alike,")
    print("  not to individual items.")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", type=Path, default=DEFAULT_DB)
    ap.add_argument("--exam", type=int, default=DEFAULT_EXAM)
    ap.add_argument("--first-pass", type=Path,
                    default=OUT_DIR / "Professor_Judgment_ID-1.xlsx")
    ap.add_argument("--agreements", type=int, default=40,
                    help="how many already-agreed items to mix in")
    ap.add_argument("--force", action="store_true",
                    help="overwrite the output even if it already holds judgments")
    ap.add_argument("--out", default="Professor_Recheck_ID-1.xlsx")
    args = ap.parse_args()
    for p in (args.db, args.first_pass):
        if not p.exists():
            sys.exit(f"not found: {p}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build(args)


if __name__ == "__main__":
    main()
