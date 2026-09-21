"""
judgment_sheet.py — the shared builder for professor judgment workbooks.

Both the first pass (build_professor_workbook.py) and the blind re-check
(build_recheck_workbook.py) produce the same thing: one row per answer, the
student's handwriting embedded as an image, a Correct/Incorrect drop-down.
Keeping the layout in one place means the two passes are visually identical,
which matters — if the re-check looked different, any change in the
professor's judgments could be blamed on the change in presentation.

What these sheets never contain: GrAid's transcription, GrAid's score, or any
trace of a previous judgment.
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

# Source resolution kept higher than the display size so zooming in Excel stays
# sharp; grayscale JPEG because these are photographs of paper, which PNG
# compresses badly (38MB of crops becomes under 3MB).
EMBED_WIDTH   = 1000
EMBED_QUALITY = 82
DISPLAY_WIDTH = 700

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
BANNER_FILL = PatternFill("solid", fgColor="FFF2CC")
DONE_FILL   = PatternFill("solid", fgColor="E2EFDA")
WRONG_FILL  = PatternFill("solid", fgColor="FCE4E4")
THIN        = Side(style="thin", color="BFBFBF")

# The task definition, worded to remove the two ways it was misread on the
# first pass: marking neat-vs-messy rather than right-vs-wrong, and accepting
# an answer that belongs to a different question.
TASK_RULES = [
    ("What counts as Correct", True),
    ("Mark Correct when the student wrote the right answer to the question shown on that row, "
     "whatever their handwriting looks like. Messy, cursive, ALL CAPS, lowercase and misspelled "
     "are all still Correct if the answer itself is right. You are judging the answer, never the "
     "penmanship.", False),
    ("", False),
    ("What counts as Incorrect", True),
    ("Mark Incorrect when the answer is wrong, left blank, or is the answer to a different "
     "question on the exam. Students sometimes write a term from elsewhere in the paper — that is "
     "Incorrect here, even though it is a real term from the subject.", False),
    ("", False),
    ("Use your own judgment on wording", True),
    ("The 'Accepted answers' column shows what the question was after. Treat it as a reference, "
     "not a rulebook. If a student writes the idea in different words, different spacing or a "
     "different spelling, and you would accept it from your own class, mark it Correct. We are "
     "measuring the gap between your judgment and the system's strict matching, so your honest "
     "opinion is exactly what we need.", False),
    ("", False),
    ("If you cannot tell", True),
    ("Leave the judgment blank and say why in the Notes column. A blank is far more useful to us "
     "than a guess.", False),
]


def _instructions(wb, title: str, intro: list[tuple[str, bool]], total: int) -> None:
    ws = wb.create_sheet("Read me first", 0)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110

    lines = [(title, True), ("", False)] + intro + [("", False)] + TASK_RULES + [
        ("", False),
        ("How to mark it", True),
        ("Click the cell in the 'Was it correct?' column and pick Correct or Incorrect from the "
         "drop-down. The row turns green or pink so you can see what you have done.", False),
        ("", False),
        (f"There are {total} answers. The Judgments sheet counts your progress at the top.", False),
    ]
    for i, (text, is_head) in enumerate(lines, start=2):
        c = ws.cell(i, 2, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if is_head:
            c.font = Font(bold=True, size=14 if i == 2 else 11,
                          color="1F3864" if i == 2 else "000000")
        if len(text) > 90:
            ws.row_dimensions[i].height = 15 * (len(text) // 90 + 1)
    ws.sheet_view.showGridLines = False


def guard_overwrite(dest: Path, force: bool) -> None:
    """
    Refuse to clobber a workbook that already holds judgments.

    Rebuilding is a normal thing to want — the accepted-answers column changes
    when the variants list is frozen. But the output filename is the same file
    the professor fills in and sends back, so an innocent rerun can destroy
    hours of somebody else's work with no warning. If the file on disk has
    verdicts in it, stop and make the caller say so explicitly.
    """
    if not dest.exists() or force:
        return
    try:
        ws = openpyxl.load_workbook(dest, data_only=True)["Judgments"]
        filled = sum(1 for r in range(5, ws.max_row + 1) if ws.cell(r, 6).value)
    except Exception:
        filled = 0
    if filled:
        raise SystemExit(
            f"{dest} already contains {filled} judgments.\n\n"
            "Rebuilding would erase them. If that is really what you want:\n"
            "  - copy the filled file somewhere safe first, then\n"
            "  - rerun with --force, or pass --out <a different name>."
        )


def build_workbook(items: list[dict], crops_dir: Path, dest: Path,
                   title: str, intro: list[tuple[str, bool]]) -> list[str]:
    """
    items: dicts with 'iid', 'prompt', 'accepted', already in final order.
    Returns the list of item ids whose crop image was missing.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Judgments"
    total = len(items)

    ws["A1"] = ("Mark each answer Correct or Incorrect. Read the 'Read me first' tab if you "
                "have not already.")
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells("A1:F1")
    ws["A1"].fill = BANNER_FILL
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24

    ws["A2"] = f'=COUNTA(F5:F{total + 4})&" of {total} answered"'
    ws["A2"].font = Font(bold=True, size=11, color="1F3864")
    ws.merge_cells("A2:C2")
    ws["D2"] = "Your work saves into this file — just keep it and send it back when you are done."
    ws["D2"].font = Font(italic=True, size=9, color="808080")
    ws.merge_cells("D2:F2")

    headers = ["#", "Item ID", "Question", "Accepted answers",
               "The student's handwriting", "Was it correct?", "Notes (optional)"]
    widths  = [5, 12, 46, 24, DISPLAY_WIDTH // 7, 16, 30]
    for col, (head, width) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(4, col, head)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[4].height = 30

    streams: list[io.BytesIO] = []   # keep alive until save()
    missing: list[str] = []

    for n, it in enumerate(items, start=1):
        row = n + 4
        iid = it["iid"]

        ws.cell(row, 1, n).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row, 2, iid).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row, 2).font = Font(size=8, color="808080")
        ws.cell(row, 3, it["prompt"]).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row, 4, it["accepted"]).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row, 4).font = Font(bold=True)
        ws.cell(row, 6).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row, 7).alignment = Alignment(wrap_text=True, vertical="center")
        for col in range(1, 8):
            ws.cell(row, col).border = Border(bottom=THIN)

        src = crops_dir / f"{iid}.png"
        if not src.exists():
            missing.append(iid)
            ws.cell(row, 5, "image missing — judge from the original paper")
            ws.row_dimensions[row].height = 30
            continue

        im = Image.open(src).convert("L")
        im = im.resize((EMBED_WIDTH, int(im.height * EMBED_WIDTH / im.width)), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=EMBED_QUALITY, optimize=True)
        buf.seek(0)
        streams.append(buf)

        pic = XLImage(buf)
        pic.width  = DISPLAY_WIDTH
        pic.height = int(DISPLAY_WIDTH * im.height / im.width)
        ws.add_image(pic, f"E{row}")
        ws.row_dimensions[row].height = (pic.height + 8) * 0.75   # px -> points

    last = total + 4
    dv = DataValidation(type="list", formula1='"Correct,Incorrect"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Pick one",
                        error="Choose Correct or Incorrect from the drop-down.")
    ws.add_data_validation(dv)
    dv.add(f"F5:F{last}")

    rng = f"A5:G{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$F5="Correct"'], fill=DONE_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$F5="Incorrect"'], fill=WRONG_FILL))

    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False

    _instructions(wb, title, intro, total)
    wb.save(dest)
    return missing
